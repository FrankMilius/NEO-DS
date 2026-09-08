import json, math, os, sys
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from lxml import etree

PX = 914400*13.3333/1920
def E(px): return Emu(int(round(px*PX)))
def rgb(h): h=h.lstrip('#'); return RGBColor(int(h[0:2],16),int(h[2:4],16),int(h[4:6],16))

FONTS = {
 'SB':('Space Grotesk',True),'SR':('Space Grotesk',False),'SMd':('Space Grotesk Medium',False),'SL':('Space Grotesk Light',False),
 'MR':('Manrope',False),'MB':('Manrope',True),'ML':('Manrope Light',False),
 'JR':('JetBrains Mono',False),'JB':('JetBrains Mono',True),
}
ALIGN = {'L':PP_ALIGN.LEFT,'C':PP_ALIGN.CENTER,'R':PP_ALIGN.RIGHT}
ANCHOR = {'T':MSO_ANCHOR.TOP,'C':MSO_ANCHOR.MIDDLE,'B':MSO_ANCHOR.BOTTOM}

recs = json.load(open('all.json'))
prs = Presentation()
prs.slide_width = Emu(int(914400*13.3333)); prs.slide_height = Emu(int(914400*7.5))
blank = prs.slide_layouts[6]

def flat(path, bg):
    im = Image.open(path).convert('RGB')
    ex = im.getextrema()
    if all(a==b for a,b in ex):
        return '#%02x%02x%02x' % (ex[0][0],ex[1][0],ex[2][0]) == bg.lower()
    return False

def paragraphs_from_segments(segs):
    """segs: [chars,font,size,color,upper] -> list of paragraphs, each list of (text,font,size,color,upper)"""
    paras=[[]]
    for chars,f,sz,col,up in segs:
        parts = chars.split('\n')
        for i,p in enumerate(parts):
            if i>0: paras.append([])
            if p: paras[-1].append((p,f,sz,col,up))
    # drop single trailing empty paragraph caused by trailing newline
    if len(paras)>1 and not paras[-1]: paras.pop()
    return paras

sections=[]  # (name, [slide_ids])
skipped_pics=0
for rec in recs:
    idx, sid, row, bg, T = rec
    s = prs.slides.add_slide(blank)
    if not sections or sections[-1][0]!=row: sections.append((row,[]))
    sections[-1][1].append(s.slide_id)
    fill = s.background.fill; fill.solid(); fill.fore_color.rgb = rgb(bg)
    png = f'png/{idx:03d}.png'
    if flat(png,bg): skipped_pics+=1
    else:
        pic = s.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height); pic.name='Bildebene'
    for tb in T:
        x,y,w,h,aH,aV,ls,lh,rot,segs = tb
        if rot:
            th=math.radians(abs(rot)); c,sn=math.cos(th),math.sin(th); d=c*c-sn*sn
            uw=(w*c-h*sn)/d; uh=(h*c-w*sn)/d
            cx,cy=x+w/2,y+h/2; w,h=uw,uh; x,y=cx-w/2,cy-h/2
        pad=max(6,0.04*w,abs(ls)+2)
        if aH=='L': w+=pad
        elif aH=='C': x-=pad/2; w+=pad
        else: x-=pad; w+=pad
        box = s.shapes.add_textbox(E(x),E(y),E(w),E(h+4))
        if rot: box.rotation = -rot
        tf = box.text_frame
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = ANCHOR[aV]
        paras = paragraphs_from_segments(segs)
        for pi,runs in enumerate(paras):
            p = tf.paragraphs[0] if pi==0 else tf.add_paragraph()
            p.alignment = ALIGN[aH]
            base_size = runs[0][2] if runs else (segs[0][2] if segs else 24)
            if lh: p.line_spacing = Pt(base_size*lh/100*0.5)
            if not runs:
                r=p.add_run(); r.text=''; r.font.size=Pt(base_size*0.5)
                fam,b=FONTS.get(segs[0][1],('Manrope',False)); r.font.name=fam
                continue
            for text,f,sz,col,up in runs:
                r = p.add_run(); r.text = text
                fam,b = FONTS.get(f,('Manrope',False))
                r.font.name=fam; r.font.bold=b; r.font.size=Pt(sz*0.5); r.font.color.rgb=rgb(col)
                rPr = r._r.get_or_add_rPr()
                if ls: rPr.set('spc', str(int(round(ls*0.5*100))))
                if up: rPr.set('cap','all')
                # set east asian / complex script font too so PowerPoint does not substitute
                for tag in ('a:latin','a:ea','a:cs'):
                    el = rPr.find(qn(tag))
                    if el is None:
                        el = etree.SubElement(rPr, qn(tag))
                    el.set('typeface', fam)
    s.notes_slide.notes_text_frame.text = f'Abschnitt: {row} · Figma-Folie {sid}'

# Sections (p14:sectionLst)
P14='http://schemas.microsoft.com/office/powerpoint/2010/main'
pres = prs.part._element
extLst = pres.find(qn('p:extLst'))
if extLst is None:
    extLst = etree.SubElement(pres, qn('p:extLst'))
ext = etree.SubElement(extLst, qn('p:ext')); ext.set('uri','{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
secLst = etree.SubElement(ext, '{%s}sectionLst'%P14, nsmap={'p14':P14})
import uuid
for name, ids in sections:
    sec = etree.SubElement(secLst, '{%s}section'%P14); sec.set('name',name); sec.set('id','{%s}'%str(uuid.uuid4()).upper())
    sl = etree.SubElement(sec, '{%s}sldIdLst'%P14)
    for i in ids:
        e = etree.SubElement(sl, '{%s}sldId'%P14); e.set('id',str(i))

out = sys.argv[1]
prs.save(out)
print('slides',len(prs.slides),'sections',len(sections),'flat pics skipped',skipped_pics, 'size MB', round(os.path.getsize(out)/1e6,1))
